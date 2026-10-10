"""Format-specific extraction. Preserve locators; never call unreadable files successful."""
import csv
import io
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field, asdict
from pathlib import Path
from defusedxml import ElementTree as ET
from PIL import Image, ImageOps


@dataclass
class Block:
    locator: str
    text: str
    kind: str = 'paragraph'
    cells: list = field(default_factory=list)


@dataclass
class Parsed:
    blocks: list[Block]
    engine: str
    warnings: list[str] = field(default_factory=list)


def decode(data):
    for enc in ('utf-8-sig', 'cp949', 'utf-16'):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise ValueError('문자 인코딩을 확인할 수 없습니다. UTF-8 텍스트로 저장해 주세요.')


def split_text(text, prefix):
    # Line ranges, not fictional page numbers. Keep paragraphs up to 1100 characters.
    blocks, buf, start = [], [], 1
    lines = text.splitlines()
    for i, line in enumerate(lines, 1):
        if buf and (sum(len(x) for x in buf) + len(line) > 1100 or not line.strip()):
            blocks.append(Block(f'{prefix} · 줄 {start}–{i-1}', '\n'.join(buf)))
            buf = []
        if line.strip():
            if not buf:
                start = i
            # Handle very long OCR lines without dropping their suffix.
            if len(line) > 1800:
                if buf:
                    blocks.append(Block(f'{prefix} · 줄 {start}–{i-1}', '\n'.join(buf)))
                    buf = []
                for offset in range(0, len(line), 1100):
                    blocks.append(Block(f'{prefix} · 줄 {i} · 구간 {offset//1100+1}', line[offset:offset+1100]))
            else:
                buf.append(line)
    if buf:
        blocks.append(Block(f'{prefix} · 줄 {start}–{len(lines)}', '\n'.join(buf)))
    return blocks


def safe_zip(path, max_size=200*1024*1024):
    z = zipfile.ZipFile(path)
    if len(z.infolist()) > 12000 or sum(i.file_size for i in z.infolist()) > max_size:
        z.close()
        raise ValueError('압축 해제 크기 또는 항목 수가 제한을 넘었습니다.')
    return z


def local_name(e):
    return e.tag.split('}')[-1]


def text_of(e, tag='t'):
    return ''.join(n.text or '' for n in e.iter() if local_name(n) == tag)


def xml_office(path, ext):
    blocks = []
    with safe_zip(path) as z:
        if ext == '.hwpx':
            names = sorted((n for n in z.namelist() if re.fullmatch(r'Contents/section\d+\.xml', n)), key=lambda n: int(re.search(r'(\d+)\.xml', n)[1]))
            if not names:
                raise ValueError('HWPX 본문 섹션이 없습니다.')
            for s, name in enumerate(names, 1):
                root = ET.fromstring(z.read(name))
                # Top-level paragraph traversal; exclude nested table paragraphs to avoid duplicates.
                def walk(e):
                    if local_name(e) == 'tbl':
                        rows = []
                        for tr in (x for x in e if local_name(x) == 'tr'):
                            rows.append([text_of(tc) for tc in tr if local_name(tc) == 'tc'])
                        for ri, row in enumerate(rows, 1):
                            blocks.append(Block(f'섹션 {s} · 표 {len(blocks)+1} · 행 {ri}', ' | '.join(row), 'table', [row]))
                        return
                    if local_name(e) == 'p':
                        parts = []
                        def para(n):
                            if local_name(n) == 'tbl':
                                walk(n)
                                return
                            if local_name(n) == 't':
                                parts.append(n.text or '')
                            for child in n:
                                para(child)
                        para(e)
                        text = ''.join(parts).strip()
                        if text:
                            blocks.extend(split_text(text, f'섹션 {s} · 문단 {len(blocks)+1}'))
                        return
                    for child in e:
                        walk(child)
                walk(root)
        elif ext == '.docx':
            root = ET.fromstring(z.read('word/document.xml'))
            body = next(e for e in root.iter() if local_name(e) == 'body')
            for i, e in enumerate(body, 1):
                if local_name(e) == 'tbl':
                    rows = [[text_of(c) for c in r if local_name(c) == 'tc'] for r in e if local_name(r) == 'tr']
                    for j, row in enumerate(rows, 1):
                        blocks.append(Block(f'표 블록 {i} · 행 {j}', ' | '.join(row), 'table', [row]))
                elif local_name(e) == 'p' and text_of(e).strip():
                    blocks.extend(split_text(text_of(e), f'문단 {i}'))
        else:
            names = sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)), key=lambda n: int(re.search(r'(\d+)\.xml', n)[1]))
            for name in names:
                slide = int(re.search(r'(\d+)\.xml', name)[1])
                root = ET.fromstring(z.read(name))
                text = '\n'.join(text_of(p) for p in root.iter() if local_name(p) == 'p')
                blocks.extend(split_text(text, f'슬라이드 {slide}'))
    return Parsed(blocks, 'HWPX XML' if ext == '.hwpx' else 'Office XML')


def ocr_image(path, settings):
    if not shutil.which('tesseract'):
        raise ValueError('OCR 엔진이 없습니다. Tesseract와 한국어 언어팩을 설치해 주세요.')
    langs = subprocess.run(['tesseract', '--list-langs'], capture_output=True, text=True, timeout=10).stdout.splitlines()[1:]
    wanted = settings.ocr_languages.split('+')
    missing = [l for l in wanted if l not in langs]
    if missing:
        raise ValueError('OCR 언어팩 미설치: ' + ', '.join(missing))
    with Image.open(path) as image:
        if image.width * image.height > 40_000_000:
            raise ValueError('이미지 해상도가 제한을 넘었습니다.')
        image = ImageOps.exif_transpose(image).convert('RGB')
        if image.width < 1800:
            factor = min(2, 1800 / image.width)
            image = image.resize((int(image.width*factor), int(image.height*factor)))
        with tempfile.TemporaryDirectory() as td:
            p = Path(td)/'ocr.png'
            image.save(p)
            r = subprocess.run(['tesseract', str(p), 'stdout', '-l', settings.ocr_languages, '--psm', '3'], capture_output=True, text=True, timeout=90)
            if r.returncode:
                raise ValueError('OCR 처리 실패. 파일 상태를 확인해 주세요.')
            text=r.stdout
            # Sparse cover letters can be mistaken for separate layout regions.
            # Retry a uniform-block layout only for short output; still warn on all OCR.
            if len(re.sub(r'\s','',text))<120:
                retry=subprocess.run(['tesseract',str(p),'stdout','-l',settings.ocr_languages,'--psm','6'],capture_output=True,text=True,timeout=90)
                if not retry.returncode and len(re.sub(r'\s','',retry.stdout))>len(re.sub(r'\s','',text)):
                    text=retry.stdout
            return text


def native_pdf(path, settings):
    from pypdf import PdfReader
    import pdfplumber
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ValueError('암호화된 PDF입니다. 암호를 해제한 파일을 사용해 주세요.')
    if len(reader.pages) > settings.max_pages:
        raise ValueError(f'PDF는 {settings.max_pages}페이지까지 처리합니다.')
    blocks, warnings = [], []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text(extraction_mode='layout') or ''
            # Per-page fallback handles mixed text/scan PDFs, not only empty whole documents.
            useful = len(re.sub(r'\s', '', text))
            broken = text.count('\ufffd') / max(1, len(text))
            if useful < 20 or broken > .03:
                try:
                    if not shutil.which('pdftoppm'):
                        raise ValueError('PDF OCR 렌더러(pdftoppm)가 없습니다.')
                    with tempfile.TemporaryDirectory() as td:
                        prefix = str(Path(td)/'page')
                        subprocess.run(['pdftoppm', '-f', str(i), '-l', str(i), '-scale-to', '2400', '-singlefile', '-png', str(path), prefix], capture_output=True, check=True, timeout=60)
                        text = ocr_image(prefix+'.png', settings)
                    warnings.append(f'{i}페이지는 OCR로 추출했습니다. 숫자·표를 원본과 대조해 주세요.')
                except Exception as e:
                    warnings.append(f'{i}페이지 인식 보완 필요: {str(e)[:240]}')
                    if useful < 20:
                        text = ''
            blocks.extend(split_text(text, f'페이지 {i}'))
            try:
                for ti, table in enumerate(pdf.pages[i-1].extract_tables(), 1):
                    for ri, row in enumerate(table, 1):
                        clean = [str(v or '').replace('\n', ' ') for v in row]
                        if any(clean):
                            blocks.append(Block(f'페이지 {i} · 표 {ti} · 행 {ri}', ' | '.join(clean), 'table', [clean]))
            except Exception:
                warnings.append(f'{i}페이지 표 구조 추출 실패. 본문 텍스트를 확인해 주세요.')
    return Parsed(blocks, 'PDF text / tables / Tesseract', warnings)


def parse(path, settings):
    path = Path(path)
    ext = path.suffix.lower()
    warnings = []
    if ext=='.pdf':
        from pypdf import PdfReader
        guard=PdfReader(path)
        if guard.is_encrypted:
            raise ValueError('암호화된 PDF입니다. 암호를 해제한 파일을 사용해 주세요.')
        if len(guard.pages)>settings.max_pages:
            raise ValueError(f'PDF는 {settings.max_pages}페이지까지 처리합니다.')
    if settings.parser == 'docling' and ext in {'.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
        import json
        try:
            r = subprocess.run([sys.executable, '-m', 'app.docling_worker', str(path), settings.ocr_languages], capture_output=True, text=True, timeout=300)
            if r.returncode:
                raise ValueError('Docling 실행 실패 또는 모델 캐시 미준비')
            data = json.loads(r.stdout)
            if not data.get('blocks') or len(data['blocks'])>settings.max_blocks:
                raise ValueError('Docling 추출 결과가 없거나 블록 제한 초과')
            return Parsed([Block(**b) for b in data['blocks']], 'Docling (local)', data.get('warnings', []))
        except Exception as e:
            warnings.append(f'Docling 처리 실패, 기본 인식기로 전환: {str(e)[:140]}')
    if ext in {'.txt', '.md'}:
        result = Parsed(split_text(decode(path.read_bytes()), '본문'), 'Text')
    elif ext in {'.csv', '.tsv'}:
        text = decode(path.read_bytes())
        rows = list(csv.reader(io.StringIO(text), delimiter='\t' if ext == '.tsv' else ','))
        header = rows[0] if rows else []
        result = Parsed([Block(f'행 {i}', ' | '.join(f'{header[j] if j<len(header) else j+1}: {v}' for j,v in enumerate(row)), 'table', [row]) for i,row in enumerate(rows,1)], 'CSV / TSV')
    elif ext in {'.hwpx', '.docx', '.pptx'}:
        result = xml_office(path, ext)
    elif ext == '.hwp':
        try:
            from hwp_hwpx_parser import Reader
        except ImportError:
            raise ValueError('HWP 파서 미설치. hwp-hwpx-parser를 설치하거나 HWPX로 저장해 주세요.')
        with Reader(str(path)) as r:
            if r.is_encrypted:
                raise ValueError('암호화된 HWP입니다.')
            blocks = split_text(r.text, 'HWP 본문')
            for ti, table in enumerate(r.tables, 1):
                for ri, row in enumerate(table.rows, 1):
                    cells = [str(x) for x in row]
                    blocks.append(Block(f'HWP 표 {ti} · 행 {ri}', ' | '.join(cells), 'table', [cells]))
            result = Parsed(blocks, 'hwp-hwpx-parser', ['HWP는 페이지 위치를 복원하지 않습니다. 본문·표 위치로 표시합니다.'])
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        with safe_zip(path):
            pass
        wb = load_workbook(path, read_only=True, data_only=False)
        cache = load_workbook(path, read_only=True, data_only=True)
        blocks = []
        try:
            for sheet in wb:
                header = []
                if sheet.sheet_state != 'visible':
                    warnings.append(f'{sheet.title}: 숨김 시트도 인식했습니다.')
                for row_number, row in enumerate(sheet.iter_rows(),1):
                    values, cells = [], [str(cell.value) if cell.value is not None else '' for cell in row]
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = str(cell.value)
                        if cell.data_type == 'f':
                            cached = cache[sheet.title][cell.coordinate].value
                            value = f'{value} (캐시값: {cached if cached is not None else "미계산"})'
                            if cached is None:
                                warnings.append(f'{sheet.title}!{cell.coordinate}: 수식 캐시가 없어 금액을 확정할 수 없습니다.')
                        values.append(f'{cell.coordinate}={value}')
                    if not values:
                        continue
                    if not header:
                        header = cells
                    blocks.append(Block(f'시트 {sheet.title} · 행 {row_number}', ('첫 비어있지 않은 행(참고): '+' | '.join(header)+'\n' if row_number>1 else '')+' | '.join(values), 'table', [cells]))
                    if len(blocks) > settings.max_blocks:
                        raise ValueError('엑셀의 행 수가 처리 제한을 넘었습니다.')
        finally:
            wb.close(); cache.close()
        result = Parsed(blocks, 'Openpyxl (cell / formula aware)')
    elif ext == '.pdf':
        result = native_pdf(path, settings)
    elif ext in {'.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
        result = Parsed(split_text(ocr_image(path, settings), '이미지 OCR'), 'Tesseract', ['OCR 결과의 날짜·금액·표는 원본 확인이 필요합니다.'])
    else:
        raise ValueError('지원하지 않는 문서 형식입니다.')
    result.warnings = list(dict.fromkeys(result.warnings + warnings))[:200]
    result.blocks = [b for b in result.blocks if b.text.strip()]
    if len(result.blocks)>settings.max_blocks:
        raise ValueError('문서의 추출 블록 수가 제한을 넘었습니다.')
    if not result.blocks:
        raise ValueError('추출한 본문이 없습니다. 스캔·암호·OCR 설정을 확인해 주세요. ' + ' / '.join(result.warnings[:3]))
    return result
