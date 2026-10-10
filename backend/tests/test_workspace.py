import io
import json
import time
import zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from app.config import Settings
from app.main import create_app
from app.parsers import parse
from app.retrieval import search, deadline_proposals


@pytest.fixture
def client(tmp_path):
    settings=Settings(data_dir=tmp_path/'data',api_token='test-workspace-key')
    with TestClient(create_app(settings)) as c:
        c.headers['X-Seomu-Key']=settings.api_token
        yield c


def wait(client,job_id):
    for _ in range(150):
        r=client.get('/api/imports/'+job_id).json()
        if r['state'] in {'completed','completed_with_issues'}:
            return r
        time.sleep(.08)
    pytest.fail('job timeout')


def upload(client,path,text):
    r=client.post('/api/imports',files=[('files',(Path(path).name,text.encode('utf-8'),'application/octet-stream'))],data={'paths':json.dumps([path])})
    assert r.status_code==200,r.text
    return wait(client,r.json()['job_id'])


def test_auth_covers_original_and_all_data(client):
    job=upload(client,'교육/안내.txt','교육 이수현황 제출기한: 2026. 10. 20.')
    doc_id=job['items'][0]['document_id']
    for path in ['/api/documents','/api/search?q=교육','/api/documents/'+doc_id+'/original','/api/tasks','/api/status']:
        assert client.get(path,headers={'X-Seomu-Key':'wrong'}).status_code==401


def test_folder_path_search_source_and_version(client):
    job=upload(client,'부서/교육/요청.txt','교육 이수현황 제출기한: 2026. 10. 20.\n수료증을 제출하여 주시기 바랍니다.')
    doc_id=job['items'][0]['document_id']
    result=client.get('/api/search',params={'q':'교육 이수현황 제출기한','folder':'부서/교육'}).json()
    assert result[0]['path']=='부서/교육/요청.txt'
    assert '줄' in result[0]['locator']
    assert not client.get('/api/search',params={'q':'교육','folder':'회계'}).json()
    unchanged=upload(client,'부서/교육/요청.txt','교육 이수현황 제출기한: 2026. 10. 20.\n수료증을 제출하여 주시기 바랍니다.')
    assert unchanged['items'][0]['state']=='unchanged'
    upload(client,'부서/교육/요청.txt','교육 이수현황 제출기한: 2026. 10. 22.')
    assert client.get('/api/documents/'+doc_id).json()['version']==2
    answer=client.post('/api/ask',json={'question':'교육 제출기한'}).json()
    assert answer['mode']=='evidence' and answer['evidence'][0]['version']==2
    assert '20.' not in answer['answer']


def test_failed_new_version_does_not_leave_stale_search(client):
    upload(client,'예산.txt','시설비 예산 검토 필요')
    r=client.post('/api/imports',files=[('files',('예산.txt',b'\xff\xfe\xfd','application/octet-stream'))],data={'paths':'["예산.txt"]'})
    wait(client,r.json()['job_id'])
    assert not client.get('/api/search',params={'q':'시설비'}).json()


def test_zip_slip_and_symlink_are_rejected_without_writes(client):
    for name,symlink in [('../outside.txt',False),('link.txt',True)]:
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as z:
            info=zipfile.ZipInfo(name)
            if symlink:
                info.external_attr=(0o120777<<16)
            z.writestr(info,'자료')
        assert client.post('/api/imports',files=[('files',('자료.zip',out.getvalue()))]).status_code==400
    assert client.get('/api/documents').json()==[]


def test_zip_keeps_subfolders_and_partial_failures(client):
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:
        z.writestr('교육/안내.txt','교육 이수현황')
        z.writestr('기타/old.xls',b'not excel')
    r=client.post('/api/imports',files=[('files',('일괄.zip',out.getvalue()))])
    job=wait(client,r.json()['job_id'])
    assert job['state']=='completed_with_issues'
    docs=client.get('/api/documents').json()
    assert {d['path'] for d in docs}=={'일괄/교육/안내.txt','일괄/기타/old.xls'}
    assert any(d['status']=='unsupported' for d in docs)


def test_duplicate_aliases_keep_paths_but_deduplicate_evidence(client):
    upload(client,'A/안내.txt','출장 정산 제출서류는 영수증입니다.')
    upload(client,'B/안내.txt','출장 정산 제출서류는 영수증입니다.')
    assert all(d['duplicate_count']==2 for d in client.get('/api/documents').json())
    assert len(client.get('/api/search',params={'q':'출장 정산'}).json())==1


def test_deadline_conflict_and_manual_confirmation(client):
    job=upload(client,'교육.txt','제출기한: 2026. 10. 20.\n붙임의 제출기한: 2026. 10. 21.')
    doc_id=job['items'][0]['document_id']
    p=client.get('/api/documents/'+doc_id+'/task-proposal').json()
    assert p['deadline']=='' and len(p['date_candidates'])==2
    body={'title':'교육 취합','deadline':'2026-10-21','source_id':doc_id,'source_version':1,'confirmed':False}
    assert client.post('/api/tasks',json=body).status_code==400
    upload(client,'교육.txt','제출기한: 2026. 10. 23.')
    body['confirmed']=True
    assert client.post('/api/tasks',json=body).status_code==400
    body['source_version']=2
    assert client.post('/api/tasks',json=body).status_code==200


def test_xlsx_preserves_cells_formulas_and_hidden_sheets(tmp_path):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title='사업비';ws.append(['사업','금액']);ws.append(['쉘터',2500]);ws['B3']='=SUM(B2:B2)'
    hidden=wb.create_sheet('검토');hidden.sheet_state='hidden';hidden.append(['확인사항','원본 대조'])
    p=tmp_path/'예산.xlsx';wb.save(p)
    result=parse(p,Settings(data_dir=tmp_path/'data',api_token='test'))
    assert any('B2=2500' in b.text and '시트 사업비' in b.locator for b in result.blocks)
    assert any('캐시' in w for w in result.warnings)
    assert any('숨김' in w for w in result.warnings)


def test_hwpx_and_docx_tables_not_flattened(tmp_path):
    hwpx=tmp_path/'요청.hwpx'
    xml='<hp:sec xmlns:hp="urn:test"><hp:p><hp:run><hp:t>자료 제출 안내</hp:t><hp:tbl><hp:tr><hp:tc><hp:p><hp:run><hp:t>성명</hp:t></hp:run></hp:p></hp:tc><hp:tc><hp:p><hp:run><hp:t>가상직원</hp:t></hp:run></hp:p></hp:tc></hp:tr></hp:tbl></hp:run></hp:p></hp:sec>'
    with zipfile.ZipFile(hwpx,'w') as z:z.writestr('Contents/section0.xml',xml)
    s=Settings(data_dir=tmp_path/'data',api_token='test')
    result=parse(hwpx,s)
    assert len(result.blocks)==2
    assert any(b.cells==[['성명','가상직원']] for b in result.blocks)
    from docx import Document
    d=Document();d.add_paragraph('서무 출장 안내');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='항목';t.cell(0,1).text='숙박비'
    p=tmp_path/'안내.docx';d.save(p)
    result=parse(p,s)
    assert any(b.cells==[['항목','숙박비']] for b in result.blocks)


def test_pptx_slide_location(tmp_path):
    from pptx import Presentation
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text='교육 자료 제출'
    path=tmp_path/'안내.pptx';p.save(path)
    result=parse(path,Settings(data_dir=tmp_path/'data',api_token='test'))
    assert any('슬라이드 1' in b.locator and '교육' in b.text for b in result.blocks)


def test_pdf_has_real_page_and_table_locators(tmp_path):
    from reportlab.pdfgen import canvas
    p=tmp_path/'manual.pdf';c=canvas.Canvas(str(p));c.drawString(70,740,'Travel expense supporting documents: receipt and report.');c.showPage();c.drawString(70,740,'Education certificates deadline: 2026-10-20.');c.save()
    result=parse(p,Settings(data_dir=tmp_path/'data',api_token='test'))
    assert any('페이지 2' in b.locator and '2026' in b.text for b in result.blocks)


def test_csv_encoding_and_formula_injection_export(client):
    upload(client,'=HYPERLINK.txt','서무 자료')
    text=client.get('/api/export/documents.csv').text
    assert "'=HYPERLINK.txt" in text


def test_local_scan_whitelist_and_symlink(client,tmp_path):
    root=tmp_path/'root';root.mkdir();(root/'업무.txt').write_text('교육 이수현황',encoding='utf-8')
    outside=tmp_path/'secret.txt';outside.write_text('외부자료');(root/'외부.txt').symlink_to(outside)
    client.app.state.service.settings.roots={'우리폴더':str(root)}
    assert client.post('/api/local-scan',json={'root':str(tmp_path)}).status_code==400
    r=client.post('/api/local-scan',json={'root':'우리폴더'})
    wait(client,r.json()['job_id'])
    assert [d['path'] for d in client.get('/api/documents').json()]==['우리폴더/업무.txt']


def test_deleting_document_removes_search_and_keeps_task(client):
    job=upload(client,'문서.txt','제출기한: 2026. 10. 20. 교육 이수자료')
    doc_id=job['items'][0]['document_id']
    client.post('/api/tasks',json={'title':'교육','source_id':doc_id,'source_version':1,'confirmed':True})
    assert client.delete('/api/documents/'+doc_id).status_code==200
    assert not client.get('/api/search',params={'q':'교육'}).json()
    assert client.get('/api/tasks').json()[0]['source_id'] is None
